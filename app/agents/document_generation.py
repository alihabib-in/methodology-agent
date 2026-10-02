"""Document generation (cross-cutting, Phase F foundation).

Renders structured artifacts to DOCX. The standardized methodology uses the
default SCAD file-output rules (A4, 1-inch margins, Arial); the SCAD-specific
methodology (Phase I) will instead apply the official template.
"""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.shared import Inches, Mm, Pt, RGBColor

from app.models.gap_assessment import GapAssessment
from app.models.scad_methodology import SCADMethodology
from app.models.standardized_methodology import StandardizedMethodology


def _configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)


def render_standardized_methodology_docx(methodology: StandardizedMethodology) -> bytes:
    doc = Document()
    _configure_document(doc)

    doc.add_heading(methodology.title, level=0)
    doc.add_paragraph(methodology.subtitle)
    doc.add_paragraph(methodology.issuing_body)
    doc.add_paragraph(f"Version: {methodology.version}")

    for section in methodology.sections:
        doc.add_heading(f"{section.number}. {section.title}", level=1)
        for line in section.content.split("\n"):
            line = line.strip()
            if not line:
                continue
            if line.startswith("- ") or line.startswith("* "):
                doc.add_paragraph(line[2:].strip(), style="List Bullet")
            else:
                doc.add_paragraph(line)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _add_annotated_paragraph(doc: Document, text: str, body_style: str | None = None) -> None:
    """Add a paragraph, rendering `[To be confirmed by SCAD]` in red italic."""
    paragraph = doc.add_paragraph(style=body_style) if body_style else doc.add_paragraph()
    parts = text.split("[To be confirmed by SCAD]")
    for index, part in enumerate(parts):
        if part:
            paragraph.add_run(part)
        if index < len(parts) - 1:
            run = paragraph.add_run("[To be confirmed by SCAD]")
            run.italic = True
            run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)


def render_scad_methodology_docx(methodology: SCADMethodology, template_path: str | None = None) -> bytes:
    doc = Document(template_path) if template_path else Document()

    style_names = {s.name for s in doc.styles}
    heading_style = "Pop_Heading1" if "Pop_Heading1" in style_names else "Heading 1"
    title_style = "Pop_Heading1" if "Pop_Heading1" in style_names else "Title"
    body_style = "SCAD_Heading1" if "SCAD_Heading1" in style_names else "Normal"

    if not template_path:
        _configure_document(doc)

    def add_heading(text: str, style: str) -> None:
        if style in style_names:
            doc.add_paragraph(text, style=style)
        else:
            paragraph = doc.add_paragraph(text)
            if paragraph.runs:
                paragraph.runs[0].bold = True

    add_heading(methodology.title, title_style)
    doc.add_paragraph(f"Version: {methodology.version}")
    doc.add_paragraph(f"Classification: {methodology.classification}")

    for section in methodology.sections:
        add_heading(f"{section.number}. {section.title}", heading_style)
        for line in section.content.split("\n"):
            line = line.strip()
            if not line:
                continue
            if line.startswith("- ") or line.startswith("* "):
                paragraph = doc.add_paragraph(style="List Bullet")
                _append_annotated_runs(paragraph, line[2:].strip())
            else:
                _add_annotated_paragraph(doc, line, body_style)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _append_annotated_runs(paragraph, text: str) -> None:
    parts = text.split("[To be confirmed by SCAD]")
    for index, part in enumerate(parts):
        if part:
            paragraph.add_run(part)
        if index < len(parts) - 1:
            run = paragraph.add_run("[To be confirmed by SCAD]")
            run.italic = True
            run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)


def review_docx(docx_bytes: bytes, expected_sections: list[str] | None = None) -> dict:
    """Basic compliance review: reopen the DOCX and report structure + flags."""
    doc = Document(BytesIO(docx_bytes))

    def is_heading(style_name: str | None) -> bool:
        if not style_name:
            return False
        return (
            style_name == "Pop_Heading1"
            or style_name == "Title"
            or style_name.startswith("Heading")
        )

    headings = [
        p.text.strip()
        for p in doc.paragraphs
        if is_heading(p.style.name if p.style else None)
    ]
    full_text = "\n".join(p.text for p in doc.paragraphs)
    unconfirmed = full_text.count("[To be confirmed by SCAD]")

    expected = expected_sections or []
    missing = [s for s in expected if s not in full_text]

    return {
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "headings": headings,
        "unconfirmed_count": unconfirmed,
        "missing_sections": missing,
    }


def render_gap_assessment_docx(gap: GapAssessment) -> bytes:
    doc = Document()
    _configure_document(doc)

    doc.add_heading("Gap Assessment Report", level=0)
    if gap.executive_summary:
        doc.add_paragraph(gap.executive_summary)

    doc.add_heading("Gap Assessment Matrix", level=1)
    if gap.matrix:
        table = doc.add_table(rows=1, cols=5)
        header = table.rows[0].cells
        for i, label in enumerate(
            ["Dimension", "International standard", "SCAD practice", "Gap level", "Description"]
        ):
            header[i].text = label
        for item in gap.matrix:
            cells = table.add_row().cells
            cells[0].text = item.dimension
            cells[1].text = item.international_standard
            cells[2].text = item.scad_practice
            cells[3].text = item.gap_level
            cells[4].text = item.description

    if gap.priority_areas:
        doc.add_heading("Priority Areas", level=1)
        for area in gap.priority_areas:
            doc.add_paragraph(
                f"{area.get('priority', '')}: {area.get('gap_description', '')} "
                f"— {area.get('recommended_action', '')}"
            )

    if gap.compliance_summary:
        doc.add_heading("Compliance Summary", level=1)
        doc.add_paragraph(gap.compliance_summary)

    if gap.recommendations:
        doc.add_heading("Recommendations", level=1)
        for recommendation in gap.recommendations:
            doc.add_paragraph(recommendation, style="List Bullet")

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
