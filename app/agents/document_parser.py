"""Word (.docx) document parsing for requirement intake (Phase D).

Uses python-docx (the lightest option — runs in-process, no separate service).
Produces a structured document with per-paragraph provenance so extracted
evidence can be traced back to a section/paragraph.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO

from docx import Document


@dataclass
class DocumentParagraph:
    index: int
    heading: str | None
    text: str


@dataclass
class ParsedDocument:
    filename: str
    title: str
    paragraphs: list[DocumentParagraph] = field(default_factory=list)

    @property
    def full_text(self) -> str:
        lines: list[str] = []
        last_heading: str | None = None
        for paragraph in self.paragraphs:
            if paragraph.heading != last_heading:
                if paragraph.heading:
                    lines.append(f"## {paragraph.heading}")
                last_heading = paragraph.heading
            lines.append(paragraph.text)
        return "\n".join(lines)

    def provenance_lines(self) -> list[dict]:
        return [
            {
                "index": p.index,
                "heading": p.heading,
                "text": p.text,
                "location": f"{p.heading or 'Document'} (paragraph {p.index})",
            }
            for p in self.paragraphs
        ]


def parse_docx(file_bytes: bytes, filename: str = "document.docx") -> ParsedDocument:
    document = Document(BytesIO(file_bytes))
    paragraphs: list[DocumentParagraph] = []
    current_heading: str | None = None
    title = ""
    index = 0

    try:
        title = (document.core_properties.title or "").strip()
    except Exception:  # noqa: BLE001 - title is best-effort
        title = ""

    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "") if para.style else ""
        if style.startswith("Heading") or style == "Title":
            current_heading = text
            if not title:
                title = text
            continue
        paragraphs.append(DocumentParagraph(index=index, heading=current_heading, text=text))
        index += 1

    return ParsedDocument(
        filename=filename,
        title=title or filename,
        paragraphs=paragraphs,
    )


def parse_pdf(file_bytes: bytes, filename: str = "document.pdf") -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(file_bytes))
    text = "\n\n".join((page.extract_text() or "") for page in reader.pages).strip()
    paragraphs = [DocumentParagraph(index=0, heading=None, text=text)] if text else []
    return ParsedDocument(filename=filename, title=filename, paragraphs=paragraphs)
