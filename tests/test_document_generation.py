from io import BytesIO

from app.agents.document_generation import render_standardized_methodology_docx, render_indicator_cards_xlsx
from app.models.indicator import IndicatorSpecification
from app.models.standardized_methodology import StandardizedMethodology, StandardizedSection


def test_render_docx_produces_valid_zip():
    methodology = StandardizedMethodology(
        title="Standardized Methodology — Test",
        sections=[
            StandardizedSection(
                number="1",
                title="Conceptual Framework and Definition",
                content="Define the indicator.\n- bullet point",
            ),
        ],
    )
    data = render_standardized_methodology_docx(methodology)
    assert isinstance(data, bytes)
    assert data[:4] == b"PK\x03\x04"  # DOCX is a ZIP archive
    assert len(data) > 1000


def test_render_indicator_cards_xlsx_transposed_layout():
    from openpyxl import load_workbook

    indicator = IndicatorSpecification(
        code="IND-001",
        topic="Economy",
        name="Population Growth Rate",
        measurement_unit="Percent",
        data_sources=["census"],
        keywords=["population", "growth"],
    )
    data = render_indicator_cards_xlsx([indicator])
    assert isinstance(data, bytes)
    assert data[:4] == b"PK\x03\x04"  # XLSX is a ZIP archive

    workbook = load_workbook(BytesIO(data))
    sheet = workbook.active
    # Column 1 holds the vertical headings in SCAD card order.
    assert sheet.cell(row=1, column=1).value == "Indicator code"
    assert sheet.cell(row=6, column=1).value == "Indicator name"
    # Column 2 holds the indicator values aligned to those headings.
    assert sheet.cell(row=1, column=2).value == "IND-001"
    assert sheet.cell(row=6, column=2).value == "Population Growth Rate"
    assert sheet.cell(row=18, column=2).value == "Percent"  # Measurement unit
    assert sheet.max_column == 2  # one indicator -> one data column
