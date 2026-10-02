from io import BytesIO

from docx import Document

from app.agents.document_parser import parse_docx


def make_docx_bytes():
    doc = Document()
    doc.core_properties.title = "AI Infrastructure Methodology"
    doc.add_heading("Objective", level=1)
    doc.add_paragraph("SCAD needs to develop a methodology for measuring AI infrastructure.")
    doc.add_heading("Requirements", level=1)
    doc.add_paragraph("The methodology must cover accelerators and data centres.")
    doc.add_paragraph("It must define statistical units.")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_parse_docx_extracts_structure():
    parsed = parse_docx(make_docx_bytes(), "ai.docx")
    assert parsed.title == "AI Infrastructure Methodology"
    assert len(parsed.paragraphs) == 3
    assert parsed.paragraphs[0].heading == "Objective"
    assert parsed.paragraphs[0].text.startswith("SCAD needs")
    assert parsed.paragraphs[1].heading == "Requirements"
    assert parsed.paragraphs[2].heading == "Requirements"
    assert "## Objective" in parsed.full_text
    assert "## Requirements" in parsed.full_text


def test_provenance_lines_have_locations():
    parsed = parse_docx(make_docx_bytes(), "ai.docx")
    lines = parsed.provenance_lines()
    assert lines[0]["location"].startswith("Objective")
